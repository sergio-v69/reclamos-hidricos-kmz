import argparse
import csv
import json
import re
import time
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl


COORD_RE = re.compile(r"(-?\d{1,3}(?:[.,]\d+)?)\s*[;,]\s*(-?\d{1,3}(?:[.,]\d+)?)")


def parse_number(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return None
    while abs(number) > 180:
        number /= 10
    return number


def in_bounds(lat, lng, bounds):
    return bounds["lat_min"] <= lat <= bounds["lat_max"] and bounds["lng_min"] <= lng <= bounds["lng_max"]


def extract_coords(*texts, bounds):
    for text in texts:
        if not text:
            continue
        for match in COORD_RE.finditer(str(text)):
            lat = parse_number(match.group(1))
            lng = parse_number(match.group(2))
            if lat is not None and lng is not None and in_bounds(lat, lng, bounds):
                return lat, lng
    return None, None


def clean_address(address, description):
    text = str(address or "").strip() or str(description or "").strip()
    text = re.sub(r"\ufeff", "", text)
    text = re.sub(r"\s+", " ", text).strip(" .,-")
    text = re.sub(r"\b(Adjunto|ABJUNTO|FOTOS?|VIDEO).*", "", text, flags=re.I).strip(" .,-")
    replacements = {
        "SARFIELD": "Velez Sarsfield",
        "velez": "Velez",
        "Mac lean": "Mac Lean",
        "RISSIONE": "Rissione",
        "asaje": "Pasaje",
    }
    for old, new in replacements.items():
        text = re.sub(old, new, text, flags=re.I)
    return text


def geocode(query, bounds):
    params = {
        "q": query,
        "format": "jsonv2",
        "limit": 1,
        "addressdetails": 0,
        "bounded": 1,
        "viewbox": f"{bounds['lng_min']},{bounds['lat_max']},{bounds['lng_max']},{bounds['lat_min']}",
    }
    url = "https://nominatim.openstreetmap.org/search?" + urllib.parse.urlencode(params)
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "reclamos-hidricos-kmz/1.0",
            "Accept": "application/json",
        },
    )
    with urllib.request.urlopen(request, timeout=20) as response:
        data = json.loads(response.read().decode("utf-8"))
    if not data:
        return None
    result = data[0]
    lat = parse_number(result.get("lat"))
    lng = parse_number(result.get("lon"))
    if lat is None or lng is None or not in_bounds(lat, lng, bounds):
        return None
    return {
        "lat": lat,
        "lng": lng,
        "display_name": result.get("display_name", ""),
        "importance": result.get("importance"),
    }


def placemark(row):
    description = (
        f"<b>Ticket:</b> {escape(str(row.get('ticket') or ''))}<br/>"
        f"<b>ID:</b> {escape(str(row.get('id') or ''))}<br/>"
        f"<b>Estado:</b> {escape(str(row.get('estado') or ''))}<br/>"
        f"<b>Problema:</b> {escape(str(row.get('problema') or ''))}<br/>"
        f"<b>Direccion:</b> {escape(str(row.get('direccion') or ''))}<br/>"
        f"<b>Fuente:</b> {escape(str(row.get('fuente') or ''))}"
    )
    name = escape(f"Ticket {row.get('ticket')} - {row.get('direccion')}"[:120])
    return f"""
    <Placemark>
      <name>{name}</name>
      <description><![CDATA[{description}]]></description>
      <Point><coordinates>{row['lng']},{row['lat']},0</coordinates></Point>
    </Placemark>"""


def find_header(headers, name):
    try:
        return headers.index(name)
    except ValueError as exc:
        available = ", ".join(str(header) for header in headers if header)
        raise SystemExit(f"No se encontro la columna '{name}'. Columnas disponibles: {available}") from exc


def load_cache(path):
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def save_cache(path, cache):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")


def process(args):
    bounds = {
        "lat_min": args.lat_min,
        "lat_max": args.lat_max,
        "lng_min": args.lng_min,
        "lng_max": args.lng_max,
    }
    input_path = Path(args.input)
    output_dir = Path(args.output_dir)
    cache_path = Path(args.cache)
    output_dir.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.load_workbook(input_path, data_only=True)
    sheet = workbook[args.sheet] if args.sheet else workbook.active
    headers = [cell.value for cell in sheet[1]]

    id_idx = find_header(headers, args.id_column)
    ticket_idx = find_header(headers, args.ticket_column)
    estado_idx = find_header(headers, args.status_column)
    address_idx = find_header(headers, args.address_column)
    problema_idx = find_header(headers, args.issue_column)
    description_idx = find_header(headers, args.description_column)
    lat_idx = find_header(headers, args.lat_column)
    lng_idx = find_header(headers, args.lng_column)

    cache = load_cache(cache_path)
    rows = []

    for excel_row in range(2, sheet.max_row + 1):
        values = [sheet.cell(excel_row, column).value for column in range(1, sheet.max_column + 1)]
        ticket = values[ticket_idx]
        address = values[address_idx]
        description = values[description_idx]
        if not ticket and not address and not description:
            continue

        lat = parse_number(values[lat_idx])
        lng = parse_number(values[lng_idx])
        source = "Lat/Lng existentes" if lat is not None and lng is not None and in_bounds(lat, lng, bounds) else ""

        if not source:
            lat, lng = extract_coords(address, description, bounds=bounds)
            source = "Coordenadas en descripcion" if lat is not None else ""

        if not source:
            normalized_address = clean_address(address, description)
            candidates = [
                f"{normalized_address}, {args.city}, {args.province}, {args.country}",
                f"{normalized_address}, {args.province}, {args.country}",
            ]
            for query in candidates:
                if not normalized_address:
                    continue
                if query not in cache:
                    try:
                        cache[query] = geocode(query, bounds)
                    except Exception as exc:
                        cache[query] = {"error": str(exc)}
                    save_cache(cache_path, cache)
                    time.sleep(args.delay)
                cached = cache.get(query)
                if isinstance(cached, dict) and "lat" in cached:
                    lat = parse_number(cached["lat"])
                    lng = parse_number(cached["lng"])
                    if lat is not None and lng is not None and in_bounds(lat, lng, bounds):
                        source = "Nominatim/OSM"
                        break

        rows.append(
            {
                "excel_row": excel_row,
                "id": values[id_idx],
                "ticket": ticket,
                "estado": values[estado_idx],
                "direccion": address,
                "problema": values[problema_idx],
                "lat": lat if source else "",
                "lng": lng if source else "",
                "fuente": source or "No resuelto",
            }
        )

    csv_path = output_dir / args.csv_name
    with csv_path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

    mapped = [row for row in rows if row["lat"] != "" and row["lng"] != ""]
    kml = f"""<?xml version="1.0" encoding="UTF-8"?>
<kml xmlns="http://www.opengis.net/kml/2.2">
  <Document>
    <name>{escape(args.name)}</name>
    <description>{len(mapped)} reclamos geolocalizados.</description>
    {''.join(placemark(row) for row in mapped)}
  </Document>
</kml>
"""
    kml_path = output_dir / "doc.kml"
    kml_path.write_text(kml, encoding="utf-8")

    kmz_path = output_dir / args.kmz_name
    with zipfile.ZipFile(kmz_path, "w", zipfile.ZIP_DEFLATED) as file:
        file.write(kml_path, "doc.kml")

    print(json.dumps({
        "total": len(rows),
        "geolocalizados": len(mapped),
        "sin_resolver": len(rows) - len(mapped),
        "kmz": str(kmz_path),
        "csv": str(csv_path),
        "cache": str(cache_path),
    }, ensure_ascii=False, indent=2))


def build_parser():
    parser = argparse.ArgumentParser(description="Geolocaliza reclamos desde una planilla y genera KMZ para Google Earth.")
    parser.add_argument("input", help="Ruta al archivo .xlsx")
    parser.add_argument("--sheet", default=None, help="Nombre de la hoja. Por defecto usa la hoja activa.")
    parser.add_argument("--output-dir", default="outputs")
    parser.add_argument("--cache", default="cache/geocode_cache.json")
    parser.add_argument("--name", default="Reclamos Hidricos geolocalizados")
    parser.add_argument("--kmz-name", default="reclamos_hidricos_geolocalizados.kmz")
    parser.add_argument("--csv-name", default="reclamos_hidricos_geolocalizados.csv")
    parser.add_argument("--city", default="Resistencia")
    parser.add_argument("--province", default="Chaco")
    parser.add_argument("--country", default="Argentina")
    parser.add_argument("--delay", type=float, default=1.1, help="Pausa entre consultas a Nominatim.")
    parser.add_argument("--lat-min", type=float, default=-27.55)
    parser.add_argument("--lat-max", type=float, default=-27.30)
    parser.add_argument("--lng-min", type=float, default=-59.10)
    parser.add_argument("--lng-max", type=float, default=-58.85)
    parser.add_argument("--id-column", default="ID")
    parser.add_argument("--ticket-column", default="Nº")
    parser.add_argument("--status-column", default="Estado")
    parser.add_argument("--address-column", default="Dirección del Ticket")
    parser.add_argument("--issue-column", default="Problema")
    parser.add_argument("--description-column", default="Descripción")
    parser.add_argument("--lat-column", default="Lat")
    parser.add_argument("--lng-column", default="Lng")
    return parser


if __name__ == "__main__":
    process(build_parser().parse_args())
